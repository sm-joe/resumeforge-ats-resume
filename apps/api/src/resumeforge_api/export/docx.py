from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from ..models import Resume


class _RichTextToDocxParser(HTMLParser):
    """
    Convert ResumeForge rich-text HTML into python-docx content.

    Supported HTML:
    - <strong>, <b>
    - <em>, <i>
    - <u>
    - <br>
    - <ul><li>
    - <ol><li>
    - plain text
    """

    def __init__(self, document: Document):
        super().__init__(convert_charrefs=True)

        self.document = document
        self.current_paragraph = None

        self.bold = False
        self.italic = False
        self.underline = False

        self.list_stack: list[str] = []

    def _ensure_paragraph(self):
        if self.current_paragraph is None:
            self.current_paragraph = self.document.add_paragraph()

        return self.current_paragraph

    def _add_text(self, text: str) -> None:
        if not text:
            return

        paragraph = self._ensure_paragraph()

        run = paragraph.add_run(text)
        run.bold = self.bold
        run.italic = self.italic
        run.underline = self.underline

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()

        if tag in {"strong", "b"}:
            self.bold = True
            return

        if tag in {"em", "i"}:
            self.italic = True
            return

        if tag == "u":
            self.underline = True
            return

        if tag == "br":
            paragraph = self._ensure_paragraph()
            paragraph.add_run().add_break()
            return

        if tag == "ul":
            self.list_stack.append("ul")
            return

        if tag == "ol":
            self.list_stack.append("ol")
            return

        if tag == "li":
            list_type = (
                self.list_stack[-1]
                if self.list_stack
                else "ul"
            )

            style = (
                "List Number"
                if list_type == "ol"
                else "List Bullet"
            )

            self.current_paragraph = self.document.add_paragraph(
                style=style
            )

    def handle_endtag(self, tag):
        tag = tag.lower()

        if tag in {"strong", "b"}:
            self.bold = False
            return

        if tag in {"em", "i"}:
            self.italic = False
            return

        if tag == "u":
            self.underline = False
            return

        if tag == "li":
            self.current_paragraph = None
            return

        if tag in {"ul", "ol"}:
            if self.list_stack:
                self.list_stack.pop()

            self.current_paragraph = None
            return

        if tag in {"p", "div"}:
            self.current_paragraph = None

    def handle_data(self, data):
        self._add_text(data)


def _add_rich_text(
    document: Document,
    value: str,
) -> None:
    """
    Add ResumeForge rich-text HTML to a Word document.
    """
    if not value or not value.strip():
        return

    parser = _RichTextToDocxParser(document)
    parser.feed(value)
    parser.close()


def resume_to_docx(
    resume: Resume,
    output_path: str | Path | None = None,
):
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(resume.profile.name or "Your Name")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(20)

    if resume.profile.headline:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(resume.profile.headline)
        run.bold = True
        run.font.size = Pt(11)

    contact_items = [
        item
        for item in (
            resume.profile.email,
            resume.profile.phone,
            resume.profile.location,
        )
        if item
    ]

    if contact_items:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(" • ".join(contact_items))
        run.font.size = Pt(9)

    links = [
        link.label or link.url
        for link in resume.profile.links
        if link.label or link.url
    ]

    if links:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(" • ".join(links))
        run.font.size = Pt(9)

    def add_heading(text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(3)

        run = paragraph.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)

    if resume.summary.strip():
        add_heading("Professional Summary")

        _add_rich_text(
            document,
            resume.summary,
        )

    if resume.experience:
        add_heading("Experience")

        for experience in resume.experience:
            paragraph = document.add_paragraph()

            run = paragraph.add_run(experience.title)
            run.bold = True

            if experience.company:
                run = paragraph.add_run(
                    f" — {experience.company}"
                )
                run.bold = True

            if experience.location:
                paragraph.add_run(
                    f" • {experience.location}"
                )

            dates = ""

            if experience.startDate:
                dates = experience.startDate

            if experience.current:
                dates = f"{dates} – Present"
            elif experience.endDate:
                dates = f"{dates} – {experience.endDate}"

            if dates:
                paragraph.add_run(
                    f" | {dates}"
                )

            for bullet in experience.bullets:
                if not bullet.strip():
                    continue

                _add_rich_text(
                    document,
                    bullet,
                )

    if resume.education:
        add_heading("Education")

        for education in resume.education:
            paragraph = document.add_paragraph()

            degree = education.degree

            if education.field:
                degree = (
                    f"{degree}, {education.field}"
                    if degree
                    else education.field
                )

            run = paragraph.add_run(
                degree or education.institution
            )
            run.bold = True

            if education.institution and degree:
                paragraph.add_run(
                    f" — {education.institution}"
                )

            if education.location:
                paragraph.add_run(
                    f" • {education.location}"
                )

            dates = ""

            if education.startDate:
                dates = education.startDate

            if education.endDate:
                dates = f"{dates} – {education.endDate}"

            if dates:
                paragraph.add_run(
                    f" | {dates}"
                )

    if resume.projects:
        add_heading("Projects")

        for project in resume.projects:
            paragraph = document.add_paragraph()

            run = paragraph.add_run(project.name)
            run.bold = True

            if project.description:
                _add_rich_text(
                    document,
                    project.description,
                )

            for bullet in project.bullets:
                if not bullet.strip():
                    continue

                _add_rich_text(
                    document,
                    bullet,
                )

            if project.url:
                document.add_paragraph(
                    project.url
                )

    if resume.skills.categories:
        add_heading("Skills")

        for category in resume.skills.categories:
            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                f"{category.name}: "
            )
            run.bold = True

            paragraph.add_run(
                ", ".join(category.items)
            )

    if resume.certifications:
        add_heading("Certifications")

        for certification in resume.certifications:
            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                certification.name
            )
            run.bold = True

            if certification.issuer:
                paragraph.add_run(
                    f" — {certification.issuer}"
                )

            if certification.date:
                paragraph.add_run(
                    f" | {certification.date}"
                )

            if certification.url:
                document.add_paragraph(
                    certification.url
                )

    if resume.languages:
        add_heading("Languages")

        for language in resume.languages:
            paragraph = document.add_paragraph()

            run = paragraph.add_run(language.name)
            run.bold = True

            if language.proficiency:
                paragraph.add_run(
                    f" — {language.proficiency}"
                )

    if resume.customSections:
        for custom_section in resume.customSections:
            if not custom_section.title:
                continue

            add_heading(custom_section.title)

            if custom_section.content:
                document.add_paragraph(
                    custom_section.content
                )

    if output_path is not None:
        output = Path(output_path)
        output.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        document.save(str(output))

        return output

    output_buffer = BytesIO()
    document.save(output_buffer)
    output_buffer.seek(0)

    return output_buffer